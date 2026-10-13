from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = r"C:\Users\ללי\PROJECTS\טריוויה\client-form.xlsx"
F = "Arial"; CYAN, CORAL, INK = "25BDF0", "F05235", "263238"
fill_in = PatternFill("solid", fgColor="FFF6C9")
fill_head = PatternFill("solid", fgColor=INK)
fill_ex = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="D5DEE3"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="center", horizontal="right")
center = Alignment(wrap_text=True, vertical="center", horizontal="center")

wb = Workbook(); ws = wb.active; ws.title = "שאלון"; ws.sheet_view.rightToLeft = True
for c, w in {"A": 6, "B": 46, "C": 26, "D": 20, "E": 20, "F": 20, "G": 30}.items():
    ws.column_dimensions[c].width = w

ws.merge_cells("A1:G1"); ws["A1"] = "🎉 שאלון טריוויה לאירוע"
ws["A1"].font = Font(name=F, size=18, bold=True, color=CORAL); ws.row_dimensions[1].height = 34
ws.merge_cells("A2:G2"); ws["A2"] = "ללי שטרן – מוזיקה בקצב · תרקדי. תשתחררי. תהני"
ws["A2"].font = Font(name=F, size=11, color=CYAN, bold=True)

legend = [
    ("איך ממלאים:", True),
    ("• ממלאים רק במשבצות הצהובות.", False),
    ("• בעמודה \"התשובה הנכונה\" כותבים את התשובה האמיתית. תשובות שגויות – לא חובה, אבל מומלץ (עד 3).", False),
    ("• שאלת נכון / לא נכון: בתשובה הנכונה כותבים \"נכון\" או \"לא נכון\", ובתשובה השגויה את ההפך.", False),
    ("• שאלת סקר (בלי תשובה נכונה): כותבים את האפשרויות בעמודות C–F, ובעמודת ההערות כותבים \"סקר\".", False),
    ("• לא צריך לסדר את התשובות – המערכת מערבבת את הסדר לבד.", False),
]
r = 4
for t, b in legend:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws.cell(r, 1, t).font = Font(name=F, size=11, bold=b, color=INK); r += 1

r += 1
fields = [("שם האירוע (יופיע על המסך):", ""), ("שם הלקוח / בעל השמחה:", ""), ("תאריך האירוע:", ""), ("ערכת עיצוב למשחק:", "DJ")]
for label, val in fields:
    ws.cell(r, 2, label).font = Font(name=F, bold=True, size=11)
    c = ws.cell(r, 3, val or None); c.fill = fill_in; c.border = box; c.font = Font(name=F)
    if "ערכת עיצוב" in label:
        dv = DataValidation(type="list", formula1='"DJ,קלאסית"', allow_blank=True); ws.add_data_validation(dv); dv.add(c)
    r += 1

r += 1
heads = ["#", "שאלה", "התשובה הנכונה", "תשובה שגויה 1", "תשובה שגויה 2", "תשובה שגויה 3", "הערות"]
for j, h in enumerate(heads, 1):
    c = ws.cell(r, j, h); c.font = Font(name=F, bold=True, color="FFFFFF", size=11)
    c.fill = fill_head; c.alignment = center; c.border = box
ws.row_dimensions[r].height = 28; header_row = r; ws.freeze_panes = ws.cell(r + 1, 1); r += 1

for ex in (["דוגמה", "באיזו עיר נולד החתן?", "חיפה", "אילת", "צפת", "נהריה", "דוגמה – לא תיכנס למשחק"],
           ["דוגמה", "מה המנה הכי טעימה הערב?", "סלט", "בשר", "קינוח", "הכול!", "סקר  (דוגמה – לא תיכנס למשחק)"]):
    for j, v in enumerate(ex, 1):
        c = ws.cell(r, j, v); c.font = Font(name=F, italic=True, color="7A7A7A", size=10)
        c.fill = fill_ex; c.border = box; c.alignment = center if j == 1 else wrap
    r += 1

for n in range(1, 31):
    for j in range(1, 8):
        c = ws.cell(r, j, n if j == 1 else None); c.border = box; c.font = Font(name=F, size=11, bold=(j == 2))
        c.alignment = center if j == 1 else wrap
        if j > 1: c.fill = fill_in
    ws.row_dimensions[r].height = 28; r += 1

r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
ws.cell(r, 1, "כשהקובץ מלא – שלחו אותו בחזרה, והוא יהפוך למשחק על המסך הגדול 🎉").font = Font(name=F, bold=True, color=CYAN, size=11)

ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = f"{header_row}:{header_row}"
wb.save(OUT); print("saved", OUT)
