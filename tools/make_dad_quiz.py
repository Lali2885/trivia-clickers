from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = r"C:\Users\ללי\Downloads\שאלון יומולדת לאבא - למילוי.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "שאלון"
ws.sheet_view.rightToLeft = True

F = "Arial"
CYAN, CORAL, INK = "25BDF0", "F05235", "263238"
fill_in = PatternFill("solid", fgColor="FFF6C9")      # צהוב = למלא
fill_done = PatternFill("solid", fgColor="EAF7FD")    # תכלת = כבר מולא
fill_head = PatternFill("solid", fgColor=INK)
fill_sec = PatternFill("solid", fgColor="FDE9E4")
fill_ex = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="D5DEE3")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="center", horizontal="right")
center = Alignment(wrap_text=True, vertical="center", horizontal="center")

widths = {"A": 5, "B": 46, "C": 26, "D": 20, "E": 20, "F": 20, "G": 28}
for c, w in widths.items():
    ws.column_dimensions[c].width = w

# ---- כותרת ----
ws.merge_cells("A1:G1")
ws["A1"] = "🎂 שאלון טריוויה – יום הולדת 57 לאבא"
ws["A1"].font = Font(name=F, size=18, bold=True, color=CORAL)
ws["A1"].alignment = Alignment(horizontal="right", vertical="center")
ws.row_dimensions[1].height = 34

ws.merge_cells("A2:G2")
ws["A2"] = "ללי שטרן – מוזיקה בקצב · תרקדי. תשתחררי. תהני"
ws["A2"].font = Font(name=F, size=11, color=CYAN, bold=True)

# ---- הוראות ----
legend = [
    "איך ממלאים:",
    "• משבצות צהובות = למלא. משבצות תכלת = כבר מולאו (אפשר לשנות אם רוצים).",
    "• בעמודה \"התשובה הנכונה\" כותבים את התשובה האמיתית.",
    "• תשובות שגויות – לא חובה. אם משאירים ריק, ללי תמציא תשובות משעשעות.",
    "• אפשר לדלג על שאלה שאין לה תשובה טובה, או לכתוב שאלה משלכם בשורות הפנויות בסוף.",
]
r = 4
for i, t in enumerate(legend):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws.cell(r, 1, t).font = Font(name=F, size=11, bold=(i == 0), color=INK)
    r += 1

# ---- פרטים כלליים ----
r += 1
ws.cell(r, 2, "השם של אבא:").font = Font(name=F, bold=True, size=11)
ws.cell(r, 3).fill = fill_in; ws.cell(r, 3).border = box
r += 1
ws.cell(r, 2, "שנת לידה (לבדיקה):").font = Font(name=F, bold=True, size=11)
c = ws.cell(r, 3, 1969); c.fill = fill_done; c.border = box; c.font = Font(name=F)
ws.cell(r, 4, "תקני אם לא נכון").font = Font(name=F, size=9, italic=True, color="6B7C85")
r += 1
ws.cell(r, 2, "ערכת עיצוב למשחק:").font = Font(name=F, bold=True, size=11)
c = ws.cell(r, 3, "DJ"); c.fill = fill_in; c.border = box; c.font = Font(name=F)
dv = DataValidation(type="list", formula1='"DJ,קלאסית"', allow_blank=True)
ws.add_data_validation(dv); dv.add(c)

# ---- טבלה ----
r += 2
headers = ["#", "שאלה", "התשובה הנכונה", "תשובה שגויה 1\n(לא חובה)", "תשובה שגויה 2\n(לא חובה)", "תשובה שגויה 3\n(לא חובה)", "הערות"]
for j, h in enumerate(headers, 1):
    c = ws.cell(r, j, h)
    c.font = Font(name=F, bold=True, color="FFFFFF", size=11)
    c.fill = fill_head; c.alignment = center; c.border = box
ws.row_dimensions[r].height = 32
ws.freeze_panes = ws.cell(r + 1, 1)
header_row = r
r += 1

# שורת דוגמה
ex = ["דוגמה", "מה הצבע האהוב על אבא?", "כחול", "ירוק", "אדום", "סגול", "שורה זו היא רק דוגמה – לא תיכנס למשחק"]
for j, v in enumerate(ex, 1):
    c = ws.cell(r, j, v)
    c.font = Font(name=F, italic=True, color="7A7A7A", size=10)
    c.fill = fill_ex; c.alignment = wrap if j > 1 else center; c.border = box
r += 1

def section(title):
    global r
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c = ws.cell(r, 1, title)
    c.font = Font(name=F, bold=True, size=12, color=CORAL)
    c.fill = fill_sec; c.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[r].height = 24
    r += 1

n = 0
def q(text, correct=None, wrongs=(None, None, None), note=None, prefilled=False):
    global r, n
    n += 1
    vals = [n, text, correct, *wrongs, note]
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v)
        c.font = Font(name=F, size=11, bold=(j == 2))
        c.border = box
        c.alignment = center if j == 1 else wrap
        if j in (3, 4, 5, 6):
            c.fill = fill_done if prefilled else fill_in
    ws.row_dimensions[r].height = 30
    r += 1

section("חלק א' – שאלות אישיות על אבא (למלא)")
for t in [
    "באיזו עיר אבא נולד?",
    "מה השם של אמא שלו (סבתא)?",
    "כמה אחים ואחיות יש לו?",
    "מה המקצוע / העבודה שלו?",
    "מה המאכל האהוב עליו?",
    "מה הוא שותה בבוקר?",
    "מה המשפט / הביטוי שהוא הכי אוהב להגיד?",
    "מה התחביב שלו?",
    "איזה שיר או זמר הוא הכי אוהב?",
    "באיזו שנה הוא התחתן?",
    "כמה ילדים יש לו?",
    "כמה נכדים יש לו?",
    "מה הדבר שהוא הכי לא סובל?",
    "איזו מכונית (או צבע מכונית) הייתה לו פעם?",
    "סיפור מצחיק או מיוחד – כתבו את השאלה ואת התשובה",
]:
    q(t)

section("חלק ב' – השנה שבה אבא נולד (כבר מולא)")
q("באיזו שנה אדם נחת לראשונה על הירח – השנה שבה אבא נולד?", "1969", ("1965", "1972", "1975"), prefilled=True)
q("מי הייתה ראש ממשלת ישראל כשאבא נולד?", "גולדה מאיר", ("דוד בן-גוריון", "לוי אשכול", "מנחם בגין"), prefilled=True)
q("נכון או לא נכון: כשאבא נולד, עוד לא היו טלפונים סלולריים", "נכון", ("לא נכון", None, None), note="שאלת נכון / לא נכון", prefilled=True)
q("כמה ימים אבא חי, בערך?", "כ-20,800", ("כ-10,000", "כ-35,000", "כ-50,000"), note="57 שנים × 365 ימים", prefilled=True)

section("חלק ג' – סקרים לאווירה (אין תשובה נכונה)")
q("מה אבא יעשה הכי הרבה ביום ההולדת?", "לאכול", ("לנאום", "להתרגש עד דמעות", "להירדם על הספה"), note="סקר – אפשר לשנות את האפשרויות", prefilled=True)
q("מה הכי מתאים לתאר את אבא? (4 תכונות)", None, (None, None, None), note="סקר – כתבו 4 תכונות בעמודות C–F")

section("שאלות נוספות משלכם (רשות)")
for _ in range(3):
    n += 1
    for j in range(1, 8):
        c = ws.cell(r, j, n if j == 1 else None)
        c.border = box; c.font = Font(name=F, size=11)
        c.alignment = center if j == 1 else wrap
        if j in (2, 3, 4, 5, 6):
            c.fill = fill_in
    ws.row_dimensions[r].height = 30
    r += 1

r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
ws.cell(r, 1, "כשהקובץ מלא – שלחו אותו בחזרה לללי, והוא יהפוך למשחק על המסך הגדול 🎉").font = Font(name=F, bold=True, color=CYAN, size=11)

ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = f"{header_row}:{header_row}"

wb.save(OUT)
print("saved", OUT, "rows", r)
